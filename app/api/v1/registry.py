from __future__ import annotations

import io
import zipfile
from datetime import datetime, timedelta, timezone
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.schemas import RegistryConfirm
from app.api.v1.admin import _complex_id
from app.domain.logic import normalize_name
from app.domain.models import (
    Account,
    AuditEvent,
    Apartment,
    Building,
    RegistryEntry,
    RegistryImport,
)
from app.security.auth import require_role
from app.web.db import get_session

router = APIRouter(prefix="/admin/registry", tags=["resident registry"])
MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024
MAX_ROWS = 20_000


def _header_key(value) -> str:
    return " ".join(str(value or "").casefold().replace("ё", "е").split())


def _parse_xlsx(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > MAX_EXPANDED:
                raise HTTPException(413, "XLSX content expands beyond the allowed size")
            if any(item.file_size > 0 and item.file_size / max(item.compress_size, 1) > 200 for item in archive.infolist()):
                raise HTTPException(413, "XLSX compression ratio is too high")
    except zipfile.BadZipFile:
        raise HTTPException(422, "Файл не является корректным XLSX") from None

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise HTTPException(413, f"В файле допускается не более {MAX_ROWS} строк")
        if sheet.max_column is not None and sheet.max_column > 50:
            raise HTTPException(422, "В файле слишком много столбцов")
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            raise HTTPException(422, "В файле нет строки заголовков")
        indexes = {_header_key(value): index for index, value in enumerate(headers)}
        name_index = next((indexes[key] for key in ("фио", "ф и о", "фамилия имя отчество", "фио жильца") if key in indexes), None)
        apartment_index = next((indexes[key] for key in ("квартира", "номер квартиры", "кв") if key in indexes), None)
        building_index = next((indexes[key] for key in ("корпус", "строение", "building") if key in indexes), None)
        if name_index is None or apartment_index is None or building_index is None:
            raise HTTPException(422, "Нужны обязательные столбцы: ФИО, Квартира, Корпус")

        rows: list[dict[str, str | int]] = []
        errors: list[dict[str, str | int]] = []
        seen: set[tuple[str, str, str]] = set()
        for line, values_row in enumerate(values, start=2):
            if not values_row or not any(value not in (None, "") for value in values_row):
                continue
            if max(name_index, apartment_index, building_index) >= len(values_row):
                errors.append({"row": line, "error": "Не заполнены ФИО, квартира или корпус"})
                continue
            full_name = str(values_row[name_index] or "").strip()
            apartment = str(values_row[apartment_index] or "").strip()
            building = str(values_row[building_index] or "").strip()
            normalized = normalize_name(full_name)
            key = (building.casefold(), apartment.casefold(), normalized)
            if not full_name or len(full_name) < 5 or not apartment or not building:
                errors.append({"row": line, "error": "Заполните ФИО, квартиру и корпус"})
            elif len(full_name) > 200 or len(apartment) > 32 or len(building) > 80:
                errors.append({"row": line, "error": "Одно из значений превышает допустимую длину"})
            elif key in seen:
                errors.append({"row": line, "error": "Дублирующаяся запись в файле"})
            else:
                seen.add(key)
                rows.append(
                    {
                        "row": line,
                        "full_name": full_name,
                        "normalized_name": normalized,
                        "apartment": apartment,
                        "building": building,
                    }
                )
        workbook.close()
        return rows, errors
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Не удалось прочитать XLSX. Сохраните файл заново и повторите") from None


@router.post("/preview", status_code=201)
async def preview_registry_import(
    file: UploadFile = File(...),
    account: Account = Depends(require_role("ADMIN")),
    session: AsyncSession = Depends(get_session),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Поддерживаются только файлы XLSX")
    data = await file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "Максимальный размер файла — 10 МБ")
    rows, errors = _parse_xlsx(data)
    if not rows and not errors:
        raise HTTPException(422, "В файле нет строк жильцов")

    complex_id = await _complex_id(session, account.max_user_id)
    existing_rows = await session.execute(
        select(RegistryEntry, Apartment, Building)
        .join(Apartment, Apartment.id == RegistryEntry.apartment_id)
        .join(Building, Building.id == Apartment.building_id)
        .where(Building.complex_id == complex_id, RegistryEntry.is_active.is_(True))
    )
    existing = {
        (building.name.casefold(), apartment.number.casefold(), entry.normalized_name): entry.full_name
        for entry, apartment, building in existing_rows
    }
    incoming = {
        (str(row["building"]).casefold(), str(row["apartment"]).casefold(), str(row["normalized_name"]))
        for row in rows
    }
    updates = [
        row
        for row in rows
        if (str(row["building"]).casefold(), str(row["apartment"]).casefold(), str(row["normalized_name"])) in existing
    ]
    new_rows = [
        row
        for row in rows
        if (str(row["building"]).casefold(), str(row["apartment"]).casefold(), str(row["normalized_name"])) not in existing
    ]
    inactive_candidates = [key for key in existing if key not in incoming]
    summary = {
        "new": len(new_rows),
        "updated": len(updates),
        "would_deactivate": len(inactive_candidates),
        "valid": len(rows),
        "errors": errors,
    }
    imported = RegistryImport(
        complex_id=complex_id,
        created_by=account.max_user_id,
        filename=file.filename[:200],
        rows_json=rows,
        summary_json=summary,
        expires_at=datetime.now(timezone.utc) + timedelta(hours=24),
    )
    session.add(imported)
    await session.commit()
    return {"import_id": str(imported.id), **summary}


@router.post("/{import_id}/confirm")
async def confirm_registry_import(
    import_id: UUID,
    body: RegistryConfirm,
    account: Account = Depends(require_role("ADMIN")),
    session: AsyncSession = Depends(get_session),
):
    complex_id = await _complex_id(session, account.max_user_id)
    imported = await session.scalar(
        select(RegistryImport).where(
            RegistryImport.id == import_id,
            RegistryImport.complex_id == complex_id,
        ).with_for_update()
    )
    if imported is None:
        raise HTTPException(404, "Предпросмотр импорта не найден")
    if imported.confirmed_at is not None:
        raise HTTPException(409, "Этот импорт уже подтверждён")
    expires_at = imported.expires_at
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    if expires_at <= datetime.now(timezone.utc):
        raise HTTPException(410, "Предпросмотр импорта истёк, загрузите файл повторно")
    if not imported.rows_json:
        raise HTTPException(422, "Нет корректных строк для импорта")
    if body.deactivate_missing and imported.summary_json.get("errors"):
        raise HTTPException(422, "Сначала исправьте строки с ошибками; отключение отсутствующих записей отменено")

    incoming_keys: set[tuple[str, str, str]] = set()
    changed = 0
    for row in imported.rows_json:
        building_name = str(row["building"])
        apartment_number = str(row["apartment"])
        normalized = str(row["normalized_name"])
        incoming_keys.add((building_name.casefold(), apartment_number.casefold(), normalized))
        building = await session.scalar(
            select(Building).where(Building.complex_id == complex_id, Building.name == building_name)
        )
        if building is None:
            building = Building(complex_id=complex_id, name=building_name)
            session.add(building)
            await session.flush()
        apartment = await session.scalar(
            select(Apartment).where(
                Apartment.building_id == building.id,
                Apartment.number == apartment_number,
            )
        )
        if apartment is None:
            apartment = Apartment(building_id=building.id, number=apartment_number)
            session.add(apartment)
            await session.flush()
        entry = await session.scalar(
            select(RegistryEntry).where(
                RegistryEntry.apartment_id == apartment.id,
                RegistryEntry.normalized_name == normalized,
            )
        )
        if entry is None:
            session.add(
                RegistryEntry(
                    apartment_id=apartment.id,
                    full_name=str(row["full_name"]),
                    normalized_name=normalized,
                    is_active=True,
                )
            )
            changed += 1
        else:
            entry.full_name = str(row["full_name"])
            entry.is_active = True
            changed += 1

    deactivated = 0
    if body.deactivate_missing:
        entries = await session.execute(
            select(RegistryEntry, Apartment, Building)
            .join(Apartment, Apartment.id == RegistryEntry.apartment_id)
            .join(Building, Building.id == Apartment.building_id)
            .where(Building.complex_id == complex_id, RegistryEntry.is_active.is_(True))
        )
        for entry, apartment, building in entries:
            key = (building.name.casefold(), apartment.number.casefold(), entry.normalized_name)
            if key not in incoming_keys:
                entry.is_active = False
                deactivated += 1
    imported.confirmed_at = datetime.now(timezone.utc)
    session.add(
        AuditEvent(
            complex_id=complex_id,
            actor_id=account.max_user_id,
            action="REGISTRY_IMPORTED",
            target_id=str(imported.id),
            metadata_json={
                "applied": changed,
                "deactivated": deactivated,
                "deactivate_missing": body.deactivate_missing,
            },
        )
    )
    # Remove staged personal data after the confirmed import has been applied.
    imported.rows_json = []
    await session.commit()
    return {"import_id": str(imported.id), "applied": changed, "deactivated": deactivated}


@router.get("")
async def registry_import_history(
    account: Account = Depends(require_role("ADMIN")),
    session: AsyncSession = Depends(get_session),
):
    complex_id = await _complex_id(session, account.max_user_id)
    imports = await session.scalars(
        select(RegistryImport)
        .where(RegistryImport.complex_id == complex_id)
        .order_by(RegistryImport.created_at.desc())
        .limit(50)
    )
    return [
        {
            "id": str(item.id),
            "filename": item.filename,
            "created_at": item.created_at.isoformat() if item.created_at else None,
            "confirmed_at": item.confirmed_at.isoformat() if item.confirmed_at else None,
            "summary": item.summary_json,
        }
        for item in imports
    ]
